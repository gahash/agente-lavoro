"""Crea data/pipeline.xlsx (se non esiste) con i fogli: Pipeline, Annunci, Collaborazioni,
Registro contatti (GDPR) e Riepilogo con formule. Non sovrascrive mai un file esistente."""
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
CFG = yaml.safe_load((ROOT / "config.yaml").read_text(encoding="utf-8"))
OUT = ROOT / CFG["pipeline"]["file"]
STATI = CFG["pipeline"]["stati"]

HEAD = PatternFill("solid", fgColor="1F2937")
FONT = Font(bold=True, color="FFFFFF")

FOGLI = {
    "Pipeline": ["ID", "Tipo (lavoro/cliente)", "Azienda", "Contatto", "Ruolo contatto", "Canale", "Stato",
                 "Primo contatto", "Ultimo contatto", "Prossimo passo", "Data prossimo passo",
                 "Valore stimato €", "Fonte", "Link", "Note"],
    "Annunci": ["ID", "Data", "Azienda", "Ruolo", "Link", "Tipo contratto", "Remoto", "RAL/tariffa",
                "Requisiti chiave", "Contatto", "Punteggio 1-10", "Motivazione", "Sospetto?", "Azione"],
    "Collaborazioni": ["ID", "Partner/cliente", "Scopo", "Condizioni", "Contratto/NDA firmato?",
                       "Titolarità codice", "Inizio", "Scadenze", "Importo €", "Pagamenti attesi",
                       "Pagato €", "Note"],
    "Registro contatti": ["ID", "Data", "Azienda", "Referente", "Fonte del dato (pubblica)",
                          "Base giuridica", "Canale", "Messaggio inviato (rif. bozza)", "Opt-out richiesto?",
                          "Data opt-out", "Note"],
}
LARGH = {"Link": 40, "Note": 40, "Motivazione": 40, "Requisiti chiave": 35, "Prossimo passo": 30,
         "Scopo": 30, "Condizioni": 30, "Messaggio inviato (rif. bozza)": 30}


def main():
    if OUT.exists():
        print(f"{OUT.name} esiste già: non lo tocco.")
        return
    wb = Workbook()
    wb.remove(wb.active)
    for nome, col in FOGLI.items():
        ws = wb.create_sheet(nome)
        ws.append(col)
        for i, c in enumerate(col, 1):
            cell = ws.cell(row=1, column=i)
            cell.fill, cell.font = HEAD, FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[cell.column_letter].width = LARGH.get(c, max(12, len(c) + 2))
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb["Pipeline"]
    dv = DataValidation(type="list", formula1='"' + ",".join(STATI) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("G2:G2000")
    dv2 = DataValidation(type="list", formula1='"lavoro,cliente,collaborazione"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("B2:B2000")
    for col in ("H", "I", "K"):
        for r in range(2, 2001):
            ws[f"{col}{r}"].number_format = "DD/MM/YYYY"
    # evidenzia prossimi passi scaduti
    ws.conditional_formatting.add("K2:K2000", CellIsRule(operator="lessThan", formula=["TODAY()"],
                                  fill=PatternFill("solid", fgColor="FECACA")))

    wa = wb["Annunci"]
    wa.conditional_formatting.add("K2:K2000", CellIsRule(operator="greaterThanOrEqual",
                                  formula=[str(CFG["ricerca_lavoro"]["soglia_contatto"])],
                                  fill=PatternFill("solid", fgColor="BBF7D0")))

    wr = wb.create_sheet("Riepilogo", 0)
    wr["A1"], wr["A1"].font = "Riepilogo pipeline", Font(bold=True, size=14)
    wr.append([])
    wr.append(["Stato", "Numero", "Valore stimato €"])
    for c in wr[3]:
        c.fill, c.font = HEAD, FONT
    for s in STATI:
        r = wr.max_row + 1
        wr.append([s, f'=COUNTIF(Pipeline!G:G,A{r})', f'=SUMIF(Pipeline!G:G,A{r},Pipeline!L:L)'])
    r = wr.max_row + 2
    wr[f"A{r}"] = "Annunci con punteggio ≥ soglia"
    wr[f"B{r}"] = f'=COUNTIF(Annunci!K:K,">={CFG["ricerca_lavoro"]["soglia_contatto"]}")'
    wr[f"A{r+1}"] = "Opt-out registrati"
    wr[f"B{r+1}"] = '=COUNTIF(\'Registro contatti\'!I:I,"si")'
    wr[f"A{r+2}"] = "Passi scaduti"
    wr[f"B{r+2}"] = '=COUNTIFS(Pipeline!K:K,"<"&TODAY(),Pipeline!K:K,"<>")'
    wr.column_dimensions["A"].width, wr.column_dimensions["B"].width, wr.column_dimensions["C"].width = 34, 12, 18

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"Creato {OUT}")


if __name__ == "__main__":
    main()
