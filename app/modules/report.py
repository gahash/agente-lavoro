"""Modulo E/F/report — pipeline, export Excel, report del mattino/sera/settimanale."""
from __future__ import annotations

from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..core import db, safety
from ..core.config import EXPORTS, REPORTS, load_settings

STATI = ["Nuovo", "Contattato", "Risposta", "Call", "Preventivo/Colloquio", "Trattativa", "Vinto", "Perso"]


def numeri() -> dict:
    s = load_settings()
    pipe = {st: db.q("SELECT COUNT(*) n FROM pipeline WHERE stato=?", (st,))[0]["n"] for st in STATI}
    return {
        "pipeline": pipe,
        "bozze_da_approvare": db.q("SELECT COUNT(*) n FROM bozze WHERE stato='da_approvare'")[0]["n"],
        "posta_da_gestire": db.q("SELECT COUNT(*) n FROM posta WHERE gestita=0 AND categoria IN ('rosso','arancio','giallo')")[0]["n"],
        "posta_allarmi": db.q("SELECT COUNT(*) n FROM posta WHERE gestita=0 AND categoria='stop'")[0]["n"],
        "opportunita_top": db.q("SELECT COUNT(*) n FROM opportunita WHERE punteggio>=? AND stato='Nuovo'", (s["soglia_annunci"],))[0]["n"],
        "aziende_nuove": db.q("SELECT COUNT(*) n FROM aziende WHERE stato IN ('Nuovo','Analizzata') AND opt_out=0")[0]["n"],
        "quota_pmi_residua": safety.pmi_quota_residua(),
        "follow_up_scaduti": db.q("SELECT COUNT(*) n FROM pipeline WHERE data_prossimo<=? AND stato='Contattato'",
                                  (date.today().isoformat(),))[0]["n"],
    }


def cruscotto(giorni: int = 1) -> dict:
    """Cosa è successo davvero nel periodo (oggi o ultimi N giorni): inviati, ricevuti, risposte, post, ricerche."""
    import re
    da = (date.today() - timedelta(days=giorni - 1)).isoformat()
    c = lambda sql, *a: db.q(sql, (da, *a))[0]["n"]
    inviate = db.q("SELECT id, tipo, destinatario, oggetto, inviato FROM bozze WHERE stato='inviata' AND inviato>=? "
                   "ORDER BY inviato DESC", (da,))
    me = (load_settings().get("email") or "").lower()
    ricevute = [m for m in db.q("SELECT id, mittente, oggetto, categoria, creato FROM posta WHERE creato>=? ORDER BY creato DESC", (da,))
                if not me or me not in (m["mittente"] or "").lower()]
    # risposte: email arrivate da un dominio a cui abbiamo scritto (in qualunque momento)
    domini = {r["destinatario"].split("@")[-1].lower()
              for r in db.q("SELECT destinatario FROM bozze WHERE stato='inviata' AND destinatario LIKE '%@%'")}
    domini -= {"gmail.com", "libero.it", "hotmail.it", "hotmail.com", "yahoo.it", "outlook.it", "icloud.com"}
    dom = lambda m: (re.findall(r"@([\w.-]+)", m or "") or [""])[-1].lower()
    risposte = [m for m in ricevute if dom(m["mittente"]) in domini]
    rimbalzi = [m for m in ricevute if re.search(r"mailer-daemon|postmaster|mail delivery|undeliver|non recapitat|delivery status",
                                                  f"{m['mittente']} {m['oggetto']}", re.I)]
    ricevute = [m for m in ricevute if m not in rimbalzi]
    from .linkedin_post import di_oggi, init as init_post
    init_post()
    prossimo = db.q("SELECT tema, data_prevista FROM post_linkedin WHERE stato='approvato' AND data_prevista>? "
                    "ORDER BY data_prevista LIMIT 1", (date.today().isoformat(),))
    return {
        "giorni": giorni,
        "inviate": len(inviate),
        "inviate_per_tipo": {t: sum(1 for b in inviate if b["tipo"] == t) for t in {b["tipo"] for b in inviate}},
        "ultime_inviate": inviate[:8],
        "ricevute": len(ricevute),
        "ultime_ricevute": ricevute[:8],
        "risposte": len(risposte),
        "rimbalzi": len(rimbalzi),
        "ultime_risposte": risposte[:5],
        "in_coda_invio": db.q("SELECT COUNT(*) n FROM bozze WHERE stato='approvata' AND destinatario LIKE '%@%'")[0]["n"],
        "a_mano": db.q("SELECT COUNT(*) n FROM bozze WHERE stato='approvata' AND (destinatario IS NULL OR destinatario NOT LIKE '%@%')")[0]["n"],
        "da_approvare": db.q("SELECT COUNT(*) n FROM bozze WHERE stato='da_approvare'")[0]["n"],
        "post_pubblicati": c("SELECT COUNT(*) n FROM post_linkedin WHERE stato='pubblicato' AND pubblicato_il>=?"),
        "post_da_pubblicare": di_oggi(),
        "prossimo_post": prossimo[0] if prossimo else None,
        "aziende_trovate": c("SELECT COUNT(*) n FROM aziende WHERE creato>=?"),
        "annunci_trovati": c("SELECT COUNT(*) n FROM opportunita WHERE creato>=?"),
        "errori": c("SELECT COUNT(*) n FROM eventi WHERE livello='errore' AND quando>=?"),
        "eventi": db.q("SELECT quando, livello, modulo, messaggio FROM eventi WHERE quando>=? AND modulo NOT IN ('app','pianificatore') "
                       "ORDER BY id DESC LIMIT 25", (da,)),
    }


def report(tipo: str = "mattino") -> dict:
    s = load_settings()
    n = numeri()
    oggi = date.today()
    righe = [f"# Report {tipo} — {oggi.strftime('%d/%m/%Y')} {datetime.now():%H:%M}", ""]
    from .attivita import riepilogo_testo
    righe += ["## Cose da fare oggi", "```", riepilogo_testo() or "(nessuna)", "```", "", "## Urgenze"]
    urg = db.q("SELECT mittente, oggetto, categoria, motivo FROM posta WHERE gestita=0 AND categoria IN ('rosso','stop') ORDER BY id DESC LIMIT 10")
    righe += [f"- {'⛔' if u['categoria']=='stop' else '🔴'} **{u['oggetto']}** — {u['mittente']} ({u['motivo']})" for u in urg] or ["- Nessuna"]
    scad = db.q("SELECT azienda, contatto, data_prossimo FROM pipeline WHERE data_prossimo<=? AND stato='Contattato' LIMIT 10",
                (oggi.isoformat(),))
    if scad:
        righe += ["", "## Follow-up da fare oggi"] + [f"- {p['azienda']} ({p['contatto']}) — previsto {p['data_prossimo']}" for p in scad]
    righe += ["", f"## Bozze da approvare: {n['bozze_da_approvare']}"]
    for b in db.q("SELECT tipo, destinatario, oggetto FROM bozze WHERE stato='da_approvare' ORDER BY id DESC LIMIT 10"):
        righe.append(f"- [{b['tipo']}] {b['oggetto']} → {b['destinatario'] or '(da completare)'}")
    righe += ["", "## Top 5 opportunità"]
    top = db.q("SELECT azienda, ruolo, punteggio, link, motivazione FROM opportunita WHERE stato='Nuovo' AND sospetto=0 "
               "ORDER BY punteggio DESC, id DESC LIMIT 5")
    righe += [f"- **{o['punteggio']}** · {o['azienda']} — {o['ruolo']} · [annuncio]({o['link']})  \n  _{o['motivazione'][:180]}_" for o in top] or ["- Nessuna"]
    righe += ["", "## Lead PMI pronti"]
    pmi = db.q("SELECT nome, settore, citta, analisi FROM aziende WHERE stato='Analizzata' AND opt_out=0 ORDER BY id DESC LIMIT 5")
    righe += [f"- {a['nome']} ({a['settore']}, {a['citta']}): {a['analisi'][:160]}" for a in pmi] or ["- Nessuno"]
    righe += ["", f"Quota contatti PMI residua oggi: {n['quota_pmi_residua']}/{s['max_pmi_giorno']}",
              "", "## Pipeline"] + [f"- {k}: {v}" for k, v in n["pipeline"].items()]
    if tipo == "settimanale":
        da = (oggi - timedelta(days=7)).isoformat()
        inviati = db.q("SELECT COUNT(*) n FROM bozze WHERE inviato>=?", (da,))[0]["n"]
        risposte = db.q("SELECT COUNT(*) n FROM pipeline WHERE stato NOT IN ('Nuovo','Contattato') AND aggiornato>=?", (da,))[0]["n"]
        tasso = f"{(risposte / inviati * 100):.0f}%" if inviati else "—"
        righe += ["", "## Settimana", f"- Messaggi inviati: {inviati}", f"- Risposte/avanzamenti: {risposte}",
                  f"- Tasso di risposta: {tasso}"]
    testo = "\n".join(righe) + "\n"
    f = REPORTS / f"{oggi.isoformat()}_{tipo}.md"
    f.write_text(testo, encoding="utf-8")
    db.log("report", f"Report {tipo} generato")
    return {"file": str(f), "testo": testo}


def esporta_excel() -> str:
    wb = Workbook()
    wb.remove(wb.active)
    head = PatternFill("solid", fgColor="1F2937")
    for nome, sql in [("Pipeline", "SELECT * FROM pipeline"), ("Annunci", "SELECT * FROM opportunita ORDER BY punteggio DESC"),
                      ("Aziende PMI", "SELECT * FROM aziende"), ("Bozze", "SELECT * FROM bozze"),
                      ("Registro contatti", "SELECT * FROM registro_contatti"), ("Posta", "SELECT * FROM posta")]:
        rows = db.q(sql)
        ws = wb.create_sheet(nome)
        cols = list(rows[0].keys()) if rows else ["(vuoto)"]
        ws.append(cols)
        for c in ws[1]:
            c.fill, c.font = head, Font(bold=True, color="FFFFFF")
        for r in rows:
            ws.append([str(v) if v is not None else "" for v in r.values()])
        ws.freeze_panes = "B2"
    f = EXPORTS / f"pipeline_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    wb.save(f)
    return str(f)
